"""
Tests for Student & Junior Researcher Express Lab features:
1. Default UI generation and preview symmetry
2. Dynamic Express Lab UI builder
3. Student Express Report Excel (.xlsx) export
"""

import os
import openpyxl
import pytest
from app import (
    get_default_student_express_lab,
    build_student_express_lab_ui,
    get_default_cohort_preview
)
from src.cohort_pipeline import ActiveIngredientCohortPipeline


def test_default_student_express_lab():
    """Verify default Student Express Lab produces non-empty, beautifully formatted HTML strings."""
    abstract_html, visuals_html, hypotheses_html = get_default_student_express_lab()
    assert isinstance(abstract_html, str) and len(abstract_html) > 100
    assert isinstance(visuals_html, str) and len(visuals_html) > 100
    assert isinstance(hypotheses_html, str) and len(hypotheses_html) > 100

    # Key pedagogical and academic sections
    assert "STUDENT &amp; JUNIOR RESEARCHER EXPRESS LAB" in abstract_html
    assert "What Was Tested?" in abstract_html
    assert "What Did Pooled Data Prove?" in abstract_html
    assert "Clinical Significance" in abstract_html
    assert "The Core Verdict" in abstract_html

    # Simplified Visuals iframe
    assert "<iframe" in visuals_html
    assert "plotly" in visuals_html.lower()

    # Hypotheses & IMRAD guide
    assert "Future Research Hypothesis" in hypotheses_html
    assert "Dose-Duration Saturation" in hypotheses_html
    assert "Combination Synergy" in hypotheses_html
    assert "Super-Responder Phenotype" in hypotheses_html
    assert "IMRAD" in hypotheses_html


def test_build_student_express_lab_ui():
    """Verify dynamic builder populates trial-specific parameters and findings."""
    mock_res = {
        "trials_count": 4,
        "total_patients": 1850,
        "meta_results": {
            "HbA1c Reduction (%)": type("MockMeta", (), {
                "re_estimate": -0.82,
                "re_ci_lower": -0.98,
                "re_ci_upper": -0.66,
                "re_p_value": 0.0001,
                "i2": 24.5,
                "tau2": 0.015,
                "measure": "%",
                "study_results": []
            })()
        },
        "verdicts": {
            "HbA1c Reduction (%)": {
                "pooled_effect": "-0.82% Mean Reduction",
                "prediction_interval": "[-1.05%, -0.59%]",
                "i2_percent": "24.5%",
                "tau2": "0.0150",
                "evidence_strength": "High-Certainty Evidence",
                "verdict_title": "Empirical Proof of Glycemic Superiority",
                "summary_narrative": "Synthesized 4 RCTs with 1,850 patients showing significant glycemic benefit."
            }
        },
        "emergent_discoveries": []
    }

    abstract_html, visuals_html, hypotheses_html = build_student_express_lab_ui("Empagliflozin", mock_res, None)

    assert "Empagliflozin" in abstract_html
    assert "HbA1c Reduction (%)" in abstract_html
    assert "1,850 Patients" in abstract_html
    assert "-0.82% Mean Reduction" in abstract_html
    assert "<iframe" in visuals_html
    assert "Empagliflozin" in hypotheses_html
    assert "IMRAD Checklist" in hypotheses_html


def test_get_default_cohort_preview_symmetry():
    """Verify get_default_cohort_preview returns exactly 25 items matching the UI wireup."""
    preview_tuple = get_default_cohort_preview()
    assert len(preview_tuple) == 25
    # The last 4 items are student express components
    cohort_student_val, student_abstract_val, student_visuals_val, student_hypotheses_val = preview_tuple[-4:]
    assert isinstance(student_abstract_val, str) and len(student_abstract_val) > 0
    assert isinstance(student_visuals_val, str) and len(student_visuals_val) > 0
    assert isinstance(student_hypotheses_val, str) and len(student_hypotheses_val) > 0


def test_student_express_excel_export():
    """Verify the Student Express Excel workbook has both dedicated sheets with correct headers."""
    pipeline = ActiveIngredientCohortPipeline("Mixed_Metabolic_Cohort")
    res = pipeline.run_all()

    student_excel_path = res.get("student_excel_path")
    assert student_excel_path is not None
    assert os.path.exists(student_excel_path)
    assert student_excel_path.endswith("_Student_Express_Report.xlsx")

    # Load openpyxl to inspect sheet structure
    wb = openpyxl.load_workbook(student_excel_path)
    sheet_names = wb.sheetnames
    assert "Executive_Student_Brief" in sheet_names
    assert "Hypotheses_and_Thesis_Guide" in sheet_names

    ws1 = wb["Executive_Student_Brief"]
    assert "MANDOROS Student & Junior Researcher Express Lab Report" in str(ws1["A1"].value)
    assert "1. 🎓 EXECUTIVE RESEARCH BRIEF" in str(ws1["A4"].value)

    ws2 = wb["Hypotheses_and_Thesis_Guide"]
    assert "Future Research Hypotheses & Thesis Blueprint" in str(ws2["A1"].value)
    assert "IMRAD CHECKLIST" in str(ws2["A11"].value) or any("IMRAD" in str(c.value) for row in ws2.iter_rows(max_row=20) for c in row if c.value)
