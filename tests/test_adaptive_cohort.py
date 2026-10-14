"""
Tests for Evidex Intelligent Adaptive Cohort Engine:
- Clinical domain mapping & standardized mean difference (SMD Hedges' g)
- Bucher anchored indirect comparisons (NMA)
- Monte Carlo SUCRA hierarchy rankings
- Extrapolation confidence scoring
- End-to-end multi-drug cohort pipeline with 8-sheet Excel and interactive HTML dashboard
"""

import sys
from pathlib import Path

# Ensure repo root is on sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
import math
import shutil
import openpyxl

from src.meta_predictive import (
    map_parameter_to_domain,
    compute_continuous_effect_size,
    compute_dichotomous_effect_size,
    compute_indirect_comparison_nma,
    compute_sucra_rankings,
    extrapolation_risk_score,
    run_subgroup_meta_analysis,
    EffectSizeResult
)
from src.cohort_pipeline import ActiveIngredientCohortPipeline


def test_domain_mapping():
    assert map_parameter_to_domain("Body Weight") == "Anthropometric & Weight Management"
    assert map_parameter_to_domain("BMI") == "Anthropometric & Weight Management"
    assert map_parameter_to_domain("HbA1c") == "Glycemic & Metabolic Regulation"
    assert map_parameter_to_domain("Fasting Glucose") == "Glycemic & Metabolic Regulation"
    assert map_parameter_to_domain("Systolic BP") == "Cardiovascular & Hemodynamic Health"


def test_hedges_g_computation():
    eff = compute_continuous_effect_size(
        study_id="Trial_1",
        arm_exp="Drug_120mg",
        arm_ctrl="Placebo",
        parameter="Body Weight",
        timepoint="12 weeks",
        n1=50,
        m1=-4.5,
        s1=2.1,
        n2=50,
        m2=-1.2,
        s2=2.0,
        measure="SMD"
    )
    assert eff is not None
    assert eff.measure == "SMD"
    assert eff.yi < 0  # Greater weight loss in experimental
    assert eff.vi > 0
    assert eff.se > 0
    assert eff.ci_lower < eff.yi < eff.ci_upper


def test_bucher_indirect_nma():
    eff_a = EffectSizeResult("S1", "Metformin", "Placebo", "Weight", "12w", "SMD", -3.0, 0.25, 0.5, -3.98, -2.02)
    eff_b = EffectSizeResult("S2", "Orlistat", "Placebo", "Weight", "12w", "SMD", -2.0, 0.25, 0.5, -2.98, -1.02)
    
    nma = compute_indirect_comparison_nma("Metformin", "Orlistat", eff_a, eff_b)
    assert nma["comparison"] == "Metformin vs. Orlistat"
    assert round(nma["indirect_effect"], 2) == -1.0
    assert round(nma["se"], 2) == round(math.sqrt(0.5), 2)
    assert nma["favored_treatment"] == "Metformin"


def test_sucra_rankings():
    estimates = {
        "Metformin": (-3.2, 0.3),
        "Orlistat": (-2.1, 0.4),
        "Placebo": (0.0, 0.2)
    }
    sucra = compute_sucra_rankings(estimates, higher_is_better=False, n_sim=500)
    assert len(sucra) == 3
    assert sucra[0]["treatment"] == "Metformin"  # Best weight reduction
    assert sucra[0]["rank"] == 1
    assert sucra[0]["sucra_score"] > sucra[1]["sucra_score"]
    assert "Gold" in sucra[0]["tier"]


def test_subgroup_meta_analysis():
    eff1 = EffectSizeResult("S1", "Metformin 500", "Placebo", "HbA1c", "12w", "MD", -0.6, 0.04, 0.2, -0.99, -0.21, dose_mg=500.0, active_ingredient="Metformin")
    eff2 = EffectSizeResult("S2", "Metformin 850", "Placebo", "HbA1c", "12w", "MD", -0.7, 0.04, 0.2, -1.09, -0.31, dose_mg=850.0, active_ingredient="Metformin")
    eff3 = EffectSizeResult("S3", "Orlistat 120", "Placebo", "HbA1c", "12w", "MD", -0.2, 0.04, 0.2, -0.59, 0.19, dose_mg=120.0, active_ingredient="Orlistat")
    eff4 = EffectSizeResult("S4", "Orlistat 120", "Placebo", "HbA1c", "12w", "MD", -0.3, 0.04, 0.2, -0.69, 0.09, dose_mg=120.0, active_ingredient="Orlistat")

    sg_res = run_subgroup_meta_analysis([eff1, eff2, eff3, eff4], subgroup_variable="active_ingredient")
    assert sg_res is not None
    assert len(sg_res.subgroups) == 2
    assert sg_res.q_between >= 0
    assert sg_res.q_between_df == 1
    assert 0 <= sg_res.p_between <= 1.0
    assert "subgroup difference" in sg_res.interpretation or "between-subgroup" in sg_res.interpretation


def test_extrapolation_confidence_levels():
    obs_drugs = ["Metformin", "Orlistat"]
    obs_doses = [500.0, 850.0, 120.0]
    obs_durs = [12.0, 52.0]

    # In bounds -> High
    r1 = extrapolation_risk_score("Metformin", 850.0, 24.0, obs_drugs, obs_doses, obs_durs)
    assert "High Confidence" in r1["level"]

    # Drug known, dose out of bounds -> Moderate
    r2 = extrapolation_risk_score("Metformin", 3500.0, 24.0, obs_drugs, obs_doses, obs_durs)
    assert "Moderate Confidence" in r2["level"]

    # Unknown drug or extreme duration -> Exploratory
    r3 = extrapolation_risk_score("Semaglutide", 2.4, 200.0, obs_drugs, obs_doses, obs_durs)
    assert "Exploratory Extrapolation" in r3["level"]


def test_mixed_metabolic_cohort_pipeline():
    cohort_dir = Path("analysis/Mixed_Metabolic_Cohort")
    if not cohort_dir.exists():
        pytest.skip("Mixed_Metabolic_Cohort directory does not exist")

    pipeline = ActiveIngredientCohortPipeline(cohort_name="Mixed_Metabolic_Cohort")
    res = pipeline.run_all()

    assert res["cohort_mode"] == "Multi-Drug Shared-Indication Cohort"
    assert "Metformin" in res["divergence_report"]["normalized_drugs"]
    assert "Orlistat" in res["divergence_report"]["normalized_drugs"]

    # Divergence Report Verification
    div = res["divergence_report"]
    assert "df_eligibility" in div
    assert not div["df_eligibility"].empty
    assert "Specifically, because" in div["adaptation_narrative"]

    # Subgroup Results Verification
    assert "subgroup_results" in res
    assert len(res["subgroup_results"]) >= 1

    # Check Master Excel
    excel_path = Path(res["excel_path"])
    assert excel_path.exists()
    wb = openpyxl.load_workbook(excel_path)
    expected_sheets = [
        'Executive_Summary', 'Cross_Trial_Comparison', 'Harmonized_Outcomes',
        'Meta_Analysis_Results', 'Predictive_Projections', 'Heterogeneity_and_Bias',
        'Trial_Differences_Report', 'Domain_Bridged_Meta_Analysis'
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames

    # Check Sheet 7 Eligibility Section and Sheet 8 Subgroup Section
    ws7 = wb["Trial_Differences_Report"]
    ws7_text = " ".join([str(c.value) for row in ws7.iter_rows() for c in row if c.value])
    assert "Eligibility Criteria Divergence" in ws7_text

    ws8 = wb["Domain_Bridged_Meta_Analysis"]
    ws8_text = " ".join([str(c.value) for row in ws8.iter_rows() for c in row if c.value])
    assert "Subgroup Meta-Analysis" in ws8_text

    # Check Interactive HTML Dashboard
    html_path = Path(res["html_path"])
    assert html_path.exists()
    with open(html_path, "r", encoding="utf-8") as f:
        html_text = f.read()
    assert "Trial Divergence & How Evidex Adapted" in html_text
    assert "Cross-Drug / Cross-Trial Domain Comparison & Ranking" in html_text
    assert "Multi-Drug Shared-Indication Cohort" in html_text

    # Check Stage-1 Packaging
    assert res.get("zip_path") is not None
    assert Path(res["zip_path"]).exists()
    assert len(res.get("individual_excels", [])) == 4
    import zipfile
    with zipfile.ZipFile(res["zip_path"], "r") as zf:
        namelist = zf.namelist()
        assert any("DPP_Metformin" in n for n in namelist)
        assert any("Orlistat" in n for n in namelist)


def test_raw_pdf_cohort_stage1_enforcement(tmp_path, monkeypatch):
    """
    Strict Pipeline Law Verification:
    Ensures that when a raw PDF is present in a cohort folder:
    1. Stage-1 extraction runs and produces <study_id>_structured_extracted.xlsx in extracted_trials/
    2. Stage-1 workbook is copied to root outputs/
    3. All individual Stage-1 workbooks are zipped into a .zip archive accessible via zip_path
    4. Individual excels are exposed in individual_excels list
    """
    from unittest.mock import MagicMock
    from src.models import (
        StudyMetadata, TreatmentArm, EligibilityCriterion,
        OutcomeMeasurement, ExtractedTrialData
    )
    import zipfile

    cohort_name = "Test_Raw_PDF_Cohort"
    cohort_dir = Path("analysis") / cohort_name
    cohort_dir.mkdir(parents=True, exist_ok=True)
    pdf_dummy = cohort_dir / "Trial_Alpha_Raw.pdf"
    pdf_dummy.write_text("Mock PDF content for testing Stage-1 extraction")

    mock_data = ExtractedTrialData(
        study=StudyMetadata(
            study_id="Trial_Alpha_Raw",
            title="Trial Alpha Random Controlled Trial",
            total_sample_size=200,
            study_phase="Phase 3"
        ),
        arms=[
            TreatmentArm(study_id="Trial_Alpha_Raw", arm_id="ARM1", arm_name="Drug A 500mg", sample_size_allocated=100, intervention_type="experimental", dose_value=500.0, dose_unit="mg"),
            TreatmentArm(study_id="Trial_Alpha_Raw", arm_id="ARM2", arm_name="Placebo", sample_size_allocated=100, intervention_type="placebo")
        ],
        eligibility=[
            EligibilityCriterion(study_id="Trial_Alpha_Raw", criterion_id="CRIT1", criterion_type="inclusion", domain="Age", criterion_name="age_min", value=18.0, unit="years", greater_than=">", equal="=")
        ],
        outcomes=[
            OutcomeMeasurement(study_id="Trial_Alpha_Raw", outcome_id="OUT1", parameter_name="HbA1c", arm_id="ARM1", timepoint_value="12", timepoint_unit="weeks", mean=-0.8, sd=0.3, sample_size=100, unit="%"),
            OutcomeMeasurement(study_id="Trial_Alpha_Raw", outcome_id="OUT2", parameter_name="HbA1c", arm_id="ARM2", timepoint_value="12", timepoint_unit="weeks", mean=-0.1, sd=0.3, sample_size=100, unit="%")
        ]
    )

    mock_extractor_cls = MagicMock()
    mock_extractor_instance = MagicMock()
    mock_extractor_instance.extract_all.return_value = mock_data
    mock_extractor_cls.return_value = mock_extractor_instance

    monkeypatch.setattr("src.extractor.RCTExtractor", mock_extractor_cls)

    pipeline = None
    try:
        pipeline = ActiveIngredientCohortPipeline(cohort_name=cohort_name, force_extract=True)
        res = pipeline.run_all()

        # 1. Verify individual Excel generated in extracted_trials/
        extracted_excel = pipeline.extracted_trials_dir / "Trial_Alpha_Raw_structured_extracted.xlsx"
        assert extracted_excel.exists()

        # 2. Verify individual Excel copied to root outputs/
        root_excel = Path("outputs") / "Trial_Alpha_Raw_structured_extracted.xlsx"
        assert root_excel.exists()

        # 3. Verify .zip bundle generated and contains the workbook
        assert res["zip_path"] is not None
        zip_path = Path(res["zip_path"])
        assert zip_path.exists()
        with zipfile.ZipFile(zip_path, "r") as zf:
            assert "Trial_Alpha_Raw_structured_extracted.xlsx" in zf.namelist()

        # 4. Verify individual_excels list
        assert len(res["individual_excels"]) >= 1
        assert any("Trial_Alpha_Raw_structured_extracted.xlsx" in p for p in res["individual_excels"])
    finally:
        shutil.rmtree(cohort_dir, ignore_errors=True)
        if pipeline and pipeline.output_dir.exists():
            shutil.rmtree(pipeline.output_dir, ignore_errors=True)
        if (Path("outputs") / "Trial_Alpha_Raw_structured_extracted.xlsx").exists():
            (Path("outputs") / "Trial_Alpha_Raw_structured_extracted.xlsx").unlink(missing_ok=True)
        if (Path("outputs") / f"{cohort_name}_Individual_Stage1_Trial_Excels.zip").exists():
            (Path("outputs") / f"{cohort_name}_Individual_Stage1_Trial_Excels.zip").unlink(missing_ok=True)

